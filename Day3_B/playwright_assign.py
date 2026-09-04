import asyncio
import os
from datetime import datetime
import json
import random

from playwright.async_api import async_playwright
from openpyxl import Workbook, load_workbook


# ============================================================
# READ CONTACTS FROM EXCEL
# ============================================================

workbook = load_workbook("contacts.xlsx")
sheet = workbook.active

headers = [cell.value for cell in sheet[1]]
print("Excel headers:", headers)

contacts = []

for row in sheet.iter_rows(min_row=2, values_only=True):

    name, phone, message = row

    # Personalize message
    if message:
        message = message.replace("{name}", str(name))

    contacts.append({
        "name": name,
        "phone": phone,
        "message": message
    })

print("Contacts loaded:", contacts)


# This list will store the result of every contact
results = []


# ============================================================
# Main Function
# ============================================================

async def main():

    async with async_playwright() as p:

        # ----------------------------------------------------
        # OPEN BROWSER
        # ----------------------------------------------------

        browser = await p.chromium.launch(headless=False)

        page = await browser.new_page()

        print("1. Browser opened")

        await page.goto("https://web.whatsapp.com")

        print("2. WhatsApp page opened")
        print("Please scan QR code if required.")

        # Wait for WhatsApp Web login
        await page.wait_for_timeout(60000)

        print("3. Wait completed")
        print("Is page closed?", page.is_closed())


        # ----------------------------------------------------
        # HANDLE WHATSAPP POPUP
        # ----------------------------------------------------

        continue_button = page.get_by_role(
            "button",
            name="Continue"
        )

        if await continue_button.count() > 0:

            print("What's new popup found.")

            await continue_button.click()

            await page.wait_for_timeout(2000)

            print("Popup closed.")

        else:

            print("No popup found.")


        # ====================================================
        # PROCESS EVERY CONTACT
        # ====================================================

        for contact in contacts:

            name = contact["name"]
            phone = contact["phone"]
            message = contact["message"]

            print()
            print("==============================")
            print(f"Processing: {name}")
            print(f"Phone: {phone}")
            print("==============================")


            # Default values for the report
            status = "failed"
            last_messages = []
            screenshot_path = ""


            try:

                # ------------------------------------------------
                # FIND SEARCH BOX
                # ------------------------------------------------

                print("4. Finding search box...")

                search_box = page.get_by_role(
                    "textbox",
                    name="Search or start a new chat"
                )

                if await search_box.count() == 0:

                    print("Search box not found.")

                    results.append({
                        "name": name,
                        "phone": phone,
                        "message": message,
                        "status": "search box not found",
                        "last_3_messages": [],
                        "screenshot": ""
                    })

                    continue


                # ------------------------------------------------
                # SEARCH CONTACT
                # ------------------------------------------------

                print("Searching for:", phone)

                await search_box.click()

                await search_box.fill(str(phone))

                # Human-like delay
                await page.wait_for_timeout(
                    random.randint(2000, 5000)
                )

                print("Search completed.")


                # ------------------------------------------------
                # FIND CONTACT
                # ------------------------------------------------

                print("Looking for contact:", name)

                result = page.get_by_text(
                    str(name),
                    exact=True
                )

                matching_results = await result.count()

                print(
                    "Matching results:",
                    matching_results
                )


                if matching_results == 0:

                    print("Contact not found.")

                    results.append({
                        "name": name,
                        "phone": phone,
                        "message": message,
                        "status": "contact not found",
                        "last_3_messages": [],
                        "screenshot": ""
                    })

                    continue


                print("Contact found.")


                # ------------------------------------------------
                # OPEN CHAT
                # ------------------------------------------------

                await result.first.click()

                await page.wait_for_timeout(
                    random.randint(2000, 5000)
                )

                print("Chat opened.")


                # ------------------------------------------------
                # FIND MESSAGE BOX
                # ------------------------------------------------

                print("Typing message...")

                message_boxes = page.get_by_role("textbox")

                textbox_count = await message_boxes.count()

                print(
                    "Number of textboxes:",
                    textbox_count
                )


                if textbox_count < 2:

                    print("Message box not found.")

                    results.append({
                        "name": name,
                        "phone": phone,
                        "message": message,
                        "status": "message box not found",
                        "last_3_messages": [],
                        "screenshot": ""
                    })

                    continue


                message_box = message_boxes.nth(1)


                # ------------------------------------------------
                # TYPE MESSAGE
                # ------------------------------------------------

                await message_box.click()

                await message_box.fill(message)

                print("Message typed.")


                # Human-like delay
                await page.wait_for_timeout(
                    random.randint(2000, 5000)
                )


                # ------------------------------------------------
                # SEND MESSAGE
                # ------------------------------------------------

                await message_box.press("Enter")

                print("Message sent.")


                # Wait after sending
                await page.wait_for_timeout(
                    random.randint(2000, 5000)
                )


                # ------------------------------------------------
                # CONFIRM SENT MESSAGE
                # ------------------------------------------------

                print("Checking sent message...")

                message_text = page.get_by_text(
                    message,
                    exact=True
                )

                sent_count = await message_text.count()

                print(
                    "Matching sent messages:",
                    sent_count
                )


                if sent_count > 0:

                    print("Sent message confirmed.")

                    status = "sent"

                else:

                    print(
                        "Could not confirm sent message."
                    )

                    status = "sent - not confirmed"


                # ------------------------------------------------
                # TAKE SCREENSHOT
                # ------------------------------------------------

                print("Taking screenshot...")

                os.makedirs(
                    "screenshots",
                    exist_ok=True
                )

                timestamp = datetime.now().strftime(
                    "%Y%m%d_%H%M%S"
                )

                screenshot_path = (
                    f"screenshots/"
                    f"{name}_{timestamp}.png"
                )

                await page.screenshot(
                    path=screenshot_path
                )

                print(
                    f"Screenshot saved: "
                    f"{screenshot_path}"
                )


                # ------------------------------------------------
                # EXTRACT LAST 3 MESSAGES
                # ------------------------------------------------

                print(
                    "Inspecting chat messages..."
                )

                message_elements = page.locator(
                    "[data-pre-plain-text]"
                )

                message_count = (
                    await message_elements.count()
                )

                print(
                    "Total message elements:",
                    message_count
                )


                start_index = max(
                    0,
                    message_count - 3
                )


                for i in range(
                    start_index,
                    message_count
                ):

                    element = message_elements.nth(i)

                    text = await element.inner_text()

                    last_messages.append(text)


                print("Last 3 messages:")

                for msg in last_messages:

                    print(msg)


                # ------------------------------------------------
                # SAVE THIS CONTACT RESULT
                # ------------------------------------------------

                results.append({
                    "name": name,
                    "phone": phone,
                    "message": message,
                    "status": status,
                    "last_3_messages": last_messages,
                    "screenshot": screenshot_path
                })


            except Exception as e:

                # ------------------------------------------------
                # ERROR HANDLING
                # ------------------------------------------------

                print(
                    f"Error processing {name}: {e}"
                )

                results.append({
                    "name": name,
                    "phone": phone,
                    "message": message,
                    "status": f"error: {str(e)}",
                    "last_3_messages": last_messages,
                    "screenshot": screenshot_path
                })


            # ------------------------------------------------
            # DELAY BEFORE NEXT CONTACT
            # ------------------------------------------------

            print(
                "Waiting before next contact..."
            )

            await page.wait_for_timeout(
                random.randint(2000, 5000)
            )


        # ====================================================
        # ALL CONTACTS FINISHED
        # ====================================================

        print()
        print("==============================")
        print("All contacts processed.")
        print("==============================")


        # ====================================================
        # CREATE JSON REPORT
        # ====================================================

        print("Creating JSON report...")

        report = {
            "date": datetime.now().strftime(
                "%Y-%m-%d"
            ),
            "contacts": results
        }


        report_filename = (
            f"whatsapp_report_"
            f"{datetime.now().strftime('%Y-%m-%d')}.json"
        )


        with open(
            report_filename,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                report,
                file,
                indent=4,
                ensure_ascii=False
            )


        print(
            f"JSON report saved: "
            f"{report_filename}"
        )


        # ====================================================
        # CREATE EXCEL REPORT
        # ====================================================

        print("Creating Excel report...")


        excel_filename = (
            f"whatsapp_report_"
            f"{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        )


        report_workbook = Workbook()

        report_sheet = (
            report_workbook.active
        )

        report_sheet.title = "WhatsApp Report"


        # ----------------------------------------------------
        # HEADER ROW
        # ----------------------------------------------------

        report_sheet.append([
            "Name",
            "Phone",
            "Message",
            "Status",
            "Last 3 Messages",
            "Screenshot"
        ])


        # ----------------------------------------------------
        # ADD ALL CONTACT RESULTS
        # ----------------------------------------------------

        for result in results:

            report_sheet.append([
                result["name"],
                result["phone"],
                result["message"],
                result["status"],
                " | ".join(
                    result["last_3_messages"]
                ),
                result["screenshot"]
            ])


        report_workbook.save(
            excel_filename
        )


        print(
            f"Excel report saved: "
            f"{excel_filename}"
        )


        # ====================================================
        # CLOSE BROWSER
        # ====================================================

        input(
            "Press Enter in the terminal "
            "to close the browser..."
        )

        await browser.close()


# ============================================================
# RUN PROGRAM
# ============================================================

asyncio.run(main())